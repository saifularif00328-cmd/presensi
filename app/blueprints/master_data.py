"""Data Master: Kelas, Siswa (+ import & kartu QR), Guru, Tahun Ajaran."""
import csv
import sqlite3
import io
import os

from flask import (Blueprint, abort, flash, redirect, render_template, request, send_file,
                   url_for)
from openpyxl import load_workbook
from PIL import Image, ImageOps

from .. import config
from ..auth import feature, roles
from ..db import execute, get_db, get_setting, new_qr_token, query
from ..qr import make_payload, qr_png
from ..services.export import kartu_pdf, xlsx
from .common import arg_int, form_int, kelas_options

bp = Blueprint("master", __name__, url_prefix="/master")

JABATAN = ["Kepala Sekolah", "Wakil Kepala Sekolah", "Guru Mapel", "Wali Kelas", "Guru Piket",
           "Guru BK", "Staf TU", "Lainnya"]


def _f(name):
    return (request.form.get(name) or "").strip() or None


# ================================================================ KELAS
@bp.route("/kelas", methods=["GET", "POST"])
@roles()
@feature("master")
def kelas():
    if request.method == "POST":
        nama = _f("nama")
        if not nama:
            flash("Nama kelas wajib diisi.", "error")
        else:
            kid = form_int("id")
            try:
                if kid:
                    execute("UPDATE kelas SET nama = ?, jenjang = ?, wali_guru_id = ? WHERE id = ?",
                            (nama, _f("jenjang"), form_int("wali_guru_id"), kid))
                    flash("Kelas diperbarui.", "success")
                else:
                    execute("INSERT INTO kelas(nama, jenjang, wali_guru_id) VALUES (?,?,?)",
                            (nama, _f("jenjang"), form_int("wali_guru_id")))
                    flash("Kelas ditambahkan.", "success")
            except sqlite3.IntegrityError:
                flash(f"Kelas '{nama}' sudah ada.", "error")
        return redirect(url_for("master.kelas"))
    rows = query("SELECT k.*, g.nama AS wali, (SELECT COUNT(*) FROM siswa s WHERE "
                 "s.kelas_id = k.id AND s.aktif = 1) AS jumlah FROM kelas k LEFT JOIN guru g "
                 "ON g.id = k.wali_guru_id ORDER BY k.jenjang, k.nama")
    edit = query("SELECT * FROM kelas WHERE id = ?", (arg_int("edit"),), one=True)
    guru = query("SELECT id, nama FROM guru WHERE aktif = 1 ORDER BY nama")
    return render_template("master/kelas.html", rows=rows, edit=edit, guru=guru)


@bp.route("/kelas/<int:kid>/hapus", methods=["POST"])
@roles()
def kelas_hapus(kid):
    n = query("SELECT COUNT(*) AS n FROM siswa WHERE kelas_id = ?", (kid,), one=True)["n"]
    if n:
        flash(f"Kelas masih memiliki {n} siswa. Pindahkan siswa terlebih dahulu.", "error")
    else:
        execute("DELETE FROM kelas WHERE id = ?", (kid,))
        flash("Kelas dihapus.", "success")
    return redirect(url_for("master.kelas"))


# ================================================================ GURU
@bp.route("/guru", methods=["GET", "POST"])
@roles()
@feature("master")
def guru():
    if request.method == "POST":
        nama = _f("nama")
        if not nama:
            flash("Nama guru wajib diisi.", "error")
            return redirect(url_for("master.guru"))
        gid = form_int("id")
        vals = (_f("nip"), nama, _f("jabatan"), _f("no_hp"), 1 if request.form.get("aktif") else 0)
        if gid:
            execute("UPDATE guru SET nip = ?, nama = ?, jabatan = ?, no_hp = ?, aktif = ? "
                    "WHERE id = ?", vals + (gid,))
            flash("Data guru diperbarui.", "success")
        else:
            execute("INSERT INTO guru(nip, nama, jabatan, no_hp, aktif) VALUES (?,?,?,?,?)", vals)
            flash("Guru ditambahkan.", "success")
        return redirect(url_for("master.guru"))
    rows = query("SELECT g.*, (SELECT GROUP_CONCAT(nama, ', ') FROM kelas WHERE "
                 "wali_guru_id = g.id) AS wali_kelas, (SELECT username FROM users u WHERE "
                 "u.guru_id = g.id LIMIT 1) AS username FROM guru g ORDER BY g.aktif DESC, g.nama")
    edit = query("SELECT * FROM guru WHERE id = ?", (arg_int("edit"),), one=True)
    return render_template("master/guru.html", rows=rows, edit=edit, jabatan=JABATAN)


@bp.route("/guru/<int:gid>/hapus", methods=["POST"])
@roles()
def guru_hapus(gid):
    execute("DELETE FROM guru WHERE id = ?", (gid,))
    flash("Data guru dihapus.", "success")
    return redirect(url_for("master.guru"))


# ================================================================ TAHUN AJARAN
@bp.route("/tahun-ajaran", methods=["GET", "POST"])
@roles()
@feature("master")
def tahun_ajaran():
    if request.method == "POST":
        nama, sem, mulai, selesai = _f("nama"), _f("semester"), _f("tanggal_mulai"), \
            _f("tanggal_selesai")
        if not (nama and sem and mulai and selesai) or mulai > selesai:
            flash("Lengkapi data dengan benar (tanggal mulai ≤ tanggal selesai).", "error")
            return redirect(url_for("master.tahun_ajaran"))
        tid = form_int("id")
        if tid:
            execute("UPDATE tahun_ajaran SET nama = ?, semester = ?, tanggal_mulai = ?, "
                    "tanggal_selesai = ? WHERE id = ?", (nama, sem, mulai, selesai, tid))
        else:
            execute("INSERT INTO tahun_ajaran(nama, semester, tanggal_mulai, tanggal_selesai) "
                    "VALUES (?,?,?,?)", (nama, sem, mulai, selesai))
        flash("Periode tahun ajaran disimpan.", "success")
        return redirect(url_for("master.tahun_ajaran"))
    rows = query("SELECT t.*, (SELECT COUNT(*) FROM presensi p WHERE p.tanggal BETWEEN "
                 "t.tanggal_mulai AND t.tanggal_selesai) AS jml_presensi FROM tahun_ajaran t "
                 "ORDER BY t.tanggal_mulai DESC")
    edit = query("SELECT * FROM tahun_ajaran WHERE id = ?", (arg_int("edit"),), one=True)
    return render_template("master/tahun_ajaran.html", rows=rows, edit=edit)


@bp.route("/tahun-ajaran/<int:tid>/aktif", methods=["POST"])
@roles()
def tahun_ajaran_aktif(tid):
    db = get_db()
    execute("UPDATE tahun_ajaran SET aktif = 0", db=db, commit=False)
    execute("UPDATE tahun_ajaran SET aktif = 1 WHERE id = ?", (tid,), db=db)
    flash("Periode aktif diganti. Periode lain tetap tersimpan sebagai arsip.", "success")
    return redirect(url_for("master.tahun_ajaran"))


@bp.route("/tahun-ajaran/<int:tid>/hapus", methods=["POST"])
@roles()
def tahun_ajaran_hapus(tid):
    execute("DELETE FROM tahun_ajaran WHERE id = ? AND aktif = 0", (tid,))
    flash("Periode dihapus (data presensi tidak ikut terhapus).", "success")
    return redirect(url_for("master.tahun_ajaran"))


# ================================================================ SISWA
@bp.route("/siswa")
@roles()
@feature("master")
def siswa():
    kelas_id = arg_int("kelas_id")
    q = request.args.get("q", "").strip()
    status = request.args.get("status", "aktif")
    where, args = [], []
    if status != "semua":
        where.append("s.aktif = ?")
        args.append(0 if status == "nonaktif" else 1)
    if kelas_id:
        where.append("s.kelas_id = ?")
        args.append(kelas_id)
    if q:
        where.append("(s.nama LIKE ? OR s.nis LIKE ? OR s.nisn LIKE ?)")
        args += [f"%{q}%"] * 3
    sql = ("SELECT s.*, k.nama AS kelas_nama FROM siswa s LEFT JOIN kelas k ON k.id = s.kelas_id"
           + (" WHERE " + " AND ".join(where) if where else "") + " ORDER BY k.nama, s.nama")
    return render_template("master/siswa.html", rows=query(sql, args), kelas=kelas_options(),
                           kelas_id=kelas_id, q=q, status=status)


def _save_foto(file, sid):
    if not file or not file.filename:
        return None
    try:
        img = Image.open(file.stream)
        img = ImageOps.exif_transpose(img).convert("RGB")
    except Exception:
        flash("File foto tidak valid (gunakan JPG/PNG).", "error")
        return None
    img.thumbnail((600, 800))
    name = f"foto/siswa_{sid}_{new_qr_token()[:6].lower()}.jpg"
    img.save(os.path.join(config.UPLOAD_DIR, name), "JPEG", quality=85)
    return name


@bp.route("/siswa/baru", methods=["GET", "POST"])
@bp.route("/siswa/<int:sid>/edit", methods=["GET", "POST"])
@roles()
@feature("master")
def siswa_form(sid=None):
    s = query("SELECT * FROM siswa WHERE id = ?", (sid,), one=True) if sid else None
    if sid and s is None:
        abort(404)
    if request.method == "POST":
        nama = _f("nama")
        if not nama:
            flash("Nama siswa wajib diisi.", "error")
            return render_template("master/siswa_form.html", s=request.form,
                                   kelas=kelas_options())
        vals = (_f("nis"), _f("nisn"), nama, _f("jk"), form_int("kelas_id"), _f("wa_ayah"),
                _f("wa_ibu"), _f("wa_wali"), 1 if request.form.get("aktif", "1") else 0)
        db = get_db()
        if s:
            execute("UPDATE siswa SET nis = ?, nisn = ?, nama = ?, jk = ?, kelas_id = ?, "
                    "wa_ayah = ?, wa_ibu = ?, wa_wali = ?, aktif = ? WHERE id = ?",
                    vals + (sid,), db=db)
        else:
            # QR code dibuat otomatis saat siswa ditambahkan
            sid = execute("INSERT INTO siswa(nis, nisn, nama, jk, kelas_id, wa_ayah, wa_ibu, "
                          "wa_wali, aktif, qr_token) VALUES (?,?,?,?,?,?,?,?,?,?)",
                          vals + (new_qr_token(),), db=db)
        foto = _save_foto(request.files.get("foto"), sid)
        if foto:
            old = s["foto"] if s else None
            execute("UPDATE siswa SET foto = ? WHERE id = ?", (foto, sid), db=db)
            if old and os.path.exists(os.path.join(config.UPLOAD_DIR, old)):
                os.remove(os.path.join(config.UPLOAD_DIR, old))
        flash("Data siswa disimpan. QR code siap dicetak.", "success")
        return redirect(url_for("main.siswa_detail", sid=sid))
    return render_template("master/siswa_form.html", s=s, kelas=kelas_options())


@bp.route("/siswa/<int:sid>/hapus", methods=["POST"])
@roles()
def siswa_hapus(sid):
    s = query("SELECT * FROM siswa WHERE id = ?", (sid,), one=True) or abort(404)
    execute("DELETE FROM siswa WHERE id = ?", (sid,))
    if s["foto"] and os.path.exists(os.path.join(config.UPLOAD_DIR, s["foto"])):
        os.remove(os.path.join(config.UPLOAD_DIR, s["foto"]))
    flash(f"Siswa {s['nama']} beserta seluruh datanya dihapus.", "success")
    return redirect(url_for("master.siswa"))


@bp.route("/siswa/<int:sid>/ganti-qr", methods=["POST"])
@roles()
def siswa_ganti_qr(sid):
    execute("UPDATE siswa SET qr_token = ? WHERE id = ?", (new_qr_token(), sid))
    flash("QR baru dibuat. Kartu lama tidak berlaku lagi.", "success")
    return redirect(url_for("main.siswa_detail", sid=sid))


@bp.route("/siswa/<int:sid>/qr.png")
@roles("piket", "bk")
def siswa_qr(sid):
    s = query("SELECT qr_token FROM siswa WHERE id = ?", (sid,), one=True) or abort(404)
    return send_file(qr_png(make_payload(s["qr_token"])), mimetype="image/png")


IMPORT_COLS = ["nis", "nisn", "nama", "jk", "kelas", "wa_ayah", "wa_ibu", "wa_wali"]


def _read_import(file):
    name = (file.filename or "").lower()
    if name.endswith(".csv"):
        text = file.stream.read().decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else None
        reader = csv.reader(io.StringIO(text), dialect) if dialect else []
        rows = [r for r in reader]
    elif name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(file.stream, read_only=True, data_only=True)
        rows = [["" if c is None else str(c) for c in r] for r in wb.active.iter_rows(values_only=True)]
    else:
        raise ValueError("Format file harus .xlsx atau .csv")
    if not rows:
        return []
    header = [h.strip().lower().replace(" ", "_") for h in rows[0]]
    if "nama" not in header:
        raise ValueError("Baris pertama harus berisi judul kolom, minimal kolom 'nama'")
    out = []
    for r in rows[1:]:
        d = {h: (r[i].strip() if i < len(r) and r[i] is not None else "") for i, h in
             enumerate(header)}
        if d.get("nama"):
            # Excel kadang mengubah NIS/nomor HP jadi angka desimal "12345.0"
            for k in ("nis", "nisn", "wa_ayah", "wa_ibu", "wa_wali"):
                if d.get(k, "").endswith(".0"):
                    d[k] = d[k][:-2]
            out.append(d)
    return out


@bp.route("/siswa/import", methods=["GET", "POST"])
@roles()
@feature("master")
def siswa_import():
    if request.method == "POST":
        file = request.files.get("file")
        if not file or not file.filename:
            flash("Pilih file terlebih dahulu.", "error")
            return redirect(url_for("master.siswa_import"))
        try:
            data = _read_import(file)
        except Exception as e:
            flash(f"Gagal membaca file: {e}", "error")
            return redirect(url_for("master.siswa_import"))
        db = get_db()
        kelas_map = {r["nama"].upper(): r["id"] for r in query("SELECT id, nama FROM kelas", db=db)}
        baru = update = 0
        for d in data:
            kid = None
            kn = d.get("kelas", "").strip()
            if kn:
                kid = kelas_map.get(kn.upper())
                if kid is None:
                    kid = execute("INSERT INTO kelas(nama) VALUES (?)", (kn,), db=db, commit=False)
                    kelas_map[kn.upper()] = kid
            vals = (d.get("nisn") or None, d["nama"], (d.get("jk") or "").upper()[:1] or None,
                    kid, d.get("wa_ayah") or None, d.get("wa_ibu") or None,
                    d.get("wa_wali") or None)
            existing = None
            if d.get("nis"):
                existing = query("SELECT id FROM siswa WHERE nis = ?", (d["nis"],), one=True, db=db)
            if existing:
                execute("UPDATE siswa SET nisn = ?, nama = ?, jk = ?, kelas_id = ?, wa_ayah = ?, "
                        "wa_ibu = ?, wa_wali = ? WHERE id = ?", vals + (existing["id"],),
                        db=db, commit=False)
                update += 1
            else:
                execute("INSERT INTO siswa(nisn, nama, jk, kelas_id, wa_ayah, wa_ibu, wa_wali, "
                        "nis, qr_token) VALUES (?,?,?,?,?,?,?,?,?)",
                        vals + (d.get("nis") or None, new_qr_token()), db=db, commit=False)
                baru += 1
        db.commit()
        flash(f"Import selesai: {baru} siswa baru, {update} diperbarui. QR code dibuat "
              "otomatis.", "success")
        return redirect(url_for("master.siswa"))
    return render_template("master/siswa_import.html", cols=IMPORT_COLS)


@bp.route("/siswa/import/template.xlsx")
@roles()
def siswa_import_template():
    buf = xlsx("Siswa", IMPORT_COLS, [
        ("1001", "0012345678", "Ahmad Fauzi", "L", "7A", "081234567890", "081298765432", ""),
        ("1002", "0012345679", "Siti Aminah", "P", "7A", "", "081211112222", ""),
    ])
    # Template import: judul kolom harus di baris pertama → tulis ulang tanpa judul besar
    wb = load_workbook(buf)
    ws = wb.active
    ws.delete_rows(1, 2)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, as_attachment=True, download_name="template-import-siswa.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ================================================================ KARTU PELAJAR
@bp.route("/kartu")
@roles()
@feature("kartu")
def kartu():
    return render_template("master/kartu.html", kelas=kelas_options(),
                           siswa=query("SELECT s.id, s.nama, k.nama AS kelas_nama FROM siswa s "
                                       "LEFT JOIN kelas k ON k.id = s.kelas_id WHERE s.aktif = 1 "
                                       "ORDER BY k.nama, s.nama"))


@bp.route("/kartu/pdf")
@roles()
@feature("kartu")
def kartu_download():
    kelas_id, siswa_id = arg_int("kelas_id"), arg_int("siswa_id")
    sql = ("SELECT s.*, k.nama AS kelas_nama FROM siswa s LEFT JOIN kelas k ON k.id = s.kelas_id "
           "WHERE s.aktif = 1")
    args = []
    if siswa_id:
        sql += " AND s.id = ?"
        args.append(siswa_id)
    elif kelas_id:
        sql += " AND s.kelas_id = ?"
        args.append(kelas_id)
    rows = query(sql + " ORDER BY k.nama, s.nama", args)
    buf = kartu_pdf(rows, get_setting("nama_sekolah"))
    name = "kartu-pelajar"
    if siswa_id and rows:
        name += f"-{rows[0]['nama'].replace(' ', '_')}"
    elif kelas_id and rows:
        name += f"-{rows[0]['kelas_nama']}"
    return send_file(buf, mimetype="application/pdf", as_attachment=False,
                     download_name=f"{name}.pdf")
