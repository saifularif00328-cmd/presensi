"""Ekspor Excel (openpyxl) dan PDF (reportlab), termasuk kartu pelajar."""
import io
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .. import config
from ..qr import make_payload, qr_png


def xlsx(title, headers, rows, subtitle=None):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    if subtitle:
        ws.append([subtitle])
    ws.append([])
    ws.append(headers)
    hdr_row = ws.max_row
    for i in range(1, len(headers) + 1):
        c = ws.cell(row=hdr_row, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2563EB")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 6), 45)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def pdf_table(title, headers, rows, subtitle=None, wide=None):
    wide = len(headers) > 7 if wide is None else wide
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4) if wide else A4,
                            leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"])]
    if subtitle:
        story.append(Paragraph(subtitle, styles["Normal"]))
    story.append(Spacer(1, 6 * mm))
    data = [headers] + [[("" if v is None else str(v)) for v in r] for r in rows]
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(t)
    doc.build(story)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- kartu pelajar
CARD_W, CARD_H = 85.6 * mm, 54 * mm  # ukuran kartu ID standar (CR80)


def _draw_card(c, x, y, s, nama_sekolah, db):
    c.setStrokeColor(colors.HexColor("#94A3B8"))
    c.setFillColor(colors.white)
    c.roundRect(x, y, CARD_W, CARD_H, 3 * mm, stroke=1, fill=1)
    # header
    c.setFillColor(colors.HexColor("#2563EB"))
    c.roundRect(x, y + CARD_H - 11 * mm, CARD_W, 11 * mm, 3 * mm, stroke=0, fill=1)
    c.rect(x, y + CARD_H - 11 * mm, CARD_W, 4 * mm, stroke=0, fill=1)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString(x + CARD_W / 2, y + CARD_H - 5 * mm, "KARTU PELAJAR")
    c.setFont("Helvetica", 7)
    c.drawCentredString(x + CARD_W / 2, y + CARD_H - 9 * mm, (nama_sekolah or "")[:48])
    # foto
    fx, fy, fw, fh = x + 4 * mm, y + 6 * mm, 22 * mm, 29 * mm
    foto = s["foto"] and os.path.join(config.UPLOAD_DIR, s["foto"])
    if foto and os.path.exists(foto):
        try:
            c.drawImage(ImageReader(foto), fx, fy, fw, fh, preserveAspectRatio=True,
                        anchor="c", mask="auto")
        except Exception:
            foto = None
    if not foto or not os.path.exists(foto):
        c.setFillColor(colors.HexColor("#E2E8F0"))
        c.rect(fx, fy, fw, fh, stroke=0, fill=1)
        c.setFillColor(colors.HexColor("#64748B"))
        c.setFont("Helvetica", 7)
        c.drawCentredString(fx + fw / 2, fy + fh / 2, "FOTO")
    # data
    c.setFillColor(colors.black)
    tx = x + 29 * mm
    c.setFont("Helvetica-Bold", 8.5)
    nama = s["nama"]
    c.drawString(tx, y + 32 * mm, nama[:26])
    if len(nama) > 26:
        c.drawString(tx, y + 28.5 * mm, nama[26:52])
    c.setFont("Helvetica", 7.5)
    c.drawString(tx, y + 24 * mm, f"NIS  : {s['nis'] or '-'}")
    c.drawString(tx, y + 20 * mm, f"NISN : {s['nisn'] or '-'}")
    c.drawString(tx, y + 16 * mm, f"Kelas: {s['kelas_nama'] or '-'}")
    # QR
    qr = ImageReader(qr_png(make_payload(s["qr_token"], db=db), box_size=6, border=1))
    q = 25 * mm
    c.drawImage(qr, x + CARD_W - q - 3 * mm, y + 4 * mm, q, q)
    c.setFont("Helvetica", 5.5)
    c.setFillColor(colors.HexColor("#64748B"))
    c.drawString(x + 4 * mm, y + 2.2 * mm, "Scan QR untuk presensi masuk & pulang")


def kartu_pdf(siswa_rows, nama_sekolah, db=None):
    """Grid 2 x 5 kartu per halaman A4, siap cetak & potong."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setTitle("Kartu Pelajar")
    pw, ph = A4
    cols, rows_per = 2, 5
    gap = 5 * mm
    mx = (pw - cols * CARD_W - (cols - 1) * gap) / 2
    my = (ph - rows_per * CARD_H - (rows_per - 1) * gap) / 2
    per_page = cols * rows_per
    for i, s in enumerate(siswa_rows):
        if i and i % per_page == 0:
            c.showPage()
        idx = i % per_page
        col, row = idx % cols, idx // cols
        x = mx + col * (CARD_W + gap)
        y = ph - my - (row + 1) * CARD_H - row * gap
        _draw_card(c, x, y, s, nama_sekolah, db)
    if not siswa_rows:
        c.drawString(40, ph - 60, "Tidak ada siswa.")
    c.showPage()
    c.save()
    buf.seek(0)
    return buf
